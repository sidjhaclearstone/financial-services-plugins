"""Build intl/Buffett_Test_Data_Europe_India.xlsx (Universe, Triage, Analysis, Sources) for the Europe & India run."""
import csv, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
from fsload import tkey
ROOT = os.path.join(os.path.dirname(__file__), '..', 'intl')
D = os.path.join(ROOT, 'data')

def f(x):
    try: return float(x)
    except (TypeError, ValueError): return None

def summ(t):
    p = os.path.join(D, tkey(t) + '_summary.json')
    return json.load(open(p)) if os.path.exists(p) else {}

HDR = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='1F3864')
def sheet(wb, name, header, rows, widths=None, wrap=()):
    ws = wb.create_sheet(name); ws.append(header)
    for c in ws[1]: c.font = HDR; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in rows: ws.append(r)
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(18, len(h) + 2)))
    for row in ws.iter_rows(min_row=2):
        for i in wrap: row[i].alignment = Alignment(wrap_text=True, vertical='top')
    return ws

wb = Workbook(); wb.remove(wb.active)
u = list(csv.DictReader(open(os.path.join(D, 'stage1_universe.csv'))))
rows = [[r['ticker'], r['name'], r['region'], r['country'], r['test1'], r['phase'], f(r['fcf_yield_used']), r['fcf_basis'], f(r['ntm_pe']), f(r['price_usd']), r['is_reit'], r['gen_ind'], r['source_note']]
        for r in u if r['test1'] in ('PASS', 'NEAR_MISS')]
rows.sort(key=lambda x: (x[4] != 'PASS', x[5] or '9', -(x[6] or 0)))
ws = sheet(wb, 'Universe', ['Ticker', 'Name', 'Region', 'Country', 'Test 1', 'Phase (1=FCF≥6%, 2=P/E<15)', 'FCF yield (USD)', 'FCF basis', 'NTM P/E', 'Price USD', 'REIT', 'FactSet gen. industry', 'Source'], rows, {'Name': 32, 'Source': 60, 'FCF basis': 22})
for row in ws.iter_rows(min_row=2): row[6].number_format = '0.0%'

tri = list(csv.DictReader(open(os.path.join(ROOT, 'triage.csv'))))
order = {'KEEP': 0, 'TOO HARD': 1, 'NO': 2}
rows = [[r['ticker'], r['name'], r['region'], r['country'], r['phase'], f(r['ntm_pe']), f(r['fcf_yield']), r['verdict'], r['reason'], r['review_override'], r['flags']] for r in tri]
rows.sort(key=lambda x: (order.get(x[7], 3), x[4], x[5] or 99))
ws = sheet(wb, 'Triage', ['Ticker', 'Name', 'Region', 'Country', 'Phase', 'NTM P/E', 'FCF yield', 'Verdict', 'Reason', 'Reviewer override', 'Flags'], rows, {'Name': 30, 'Reason': 90, 'Reviewer override': 60}, wrap=(8, 9))
for row in ws.iter_rows(min_row=2): row[6].number_format = '0.0%'

res = {}
for line in open(os.path.join(ROOT, 'analyses', '_results_final.jsonl')):
    if line.strip(): r = json.loads(line); res[r['ticker']] = r
cols = ['ticker', 'name', 'region', 'country', 'phase', 'currency', 'ntm_pe', 'fcf_yield', 'fcf_check', 'mkt_cap_bn', 'eps_cagr_10y', 'cons_growth_fy1', 'cons_growth_fy2', 'cons_growth_fy3',
        'driver_sum_pct', 'top_drivers', 'test1', 'p_test2', 'test2', 'p_test3', 'test3', 'tests_passed', 'near_miss', 'specialness', 'confidence', 'earnings_quality_flag', 'key_uncertainty', 'calibration_note']
rows = []
for t, r in res.items():
    s = summ(t)
    rows.append([r.get(c) for c in cols] + [s.get('years_eps_growth_ge7'), s.get('fcf_conversion_10y'), s.get('net_debt_ebitda'), s.get('roic_10y_avg'), s.get('roe_10y_avg')])
rows.sort(key=lambda x: (-(x[21] or 0), -(x[17] or 0)))
ws = sheet(wb, 'Analysis', [c.replace('_', ' ') for c in cols] + ['yrs EPS growth >=7% (of 10)', 'FCF conv 10y', 'Net debt/EBITDA', 'ROIC 10y avg', 'ROE 10y avg'], rows,
           {'name': 30, 'top drivers': 40, 'specialness': 40, 'key uncertainty': 70, 'earnings quality flag': 40, 'calibration note': 50})
for row in ws.iter_rows(min_row=2):
    for i in (17, 19):
        if isinstance(row[i].value, (int, float)): row[i].number_format = '0%'

src = [
    ['Universe', 'FactSet CompanyScreener (USD)', 'HQ country in Europe list (EU/EEA + UK, CH, NO, IS, JE, GG, IM; excl. RU, BY, UA, TR) or India; ISON_SECONDARY=0; FF_MKT_VAL>$1bn. Prefilters: FY1 P/FCF 0–25, FF_FCF_YLD≥4, FY2 P/E 0–17', '2026-10-19'],
    ['Price', 'FactSet GlobalPrices close (USD)', 'latest on/before 19-Oct-2026', '2026-10-19'],
    ['NTM EPS / FCFPS / FFO', 'FactSet EstimatesConsensus NTMA (USD)', 'FCF yield = NTM FCFPS/price; fallback LTM (FF_OPER_CF − |FF_CAPEX|)/(diluted shares × price)', '2026-10-19'],
    ['FY0–FY3 consensus, LTG', 'FactSet EstimatesConsensus consensus_rolling ANN, currency ESTIMATE (reporting/broker currency)', '', '2026-10-19'],
    ['10-yr fundamentals', 'FactSet Fundamentals ANN, currency DOC (reporting currency), FY2015–FY2026', 'audit-card formula stored per value in intl/data/<TICKER>.json', 'pulled 2026-10-20'],
    ['Adjusted EPS history', 'FactSet EstimatesConsensus surprise actuals (ESTIMATE currency)', '', 'pulled 2026-10-20'],
    ['Business description, RBICS', 'FactSet EntityReference', '', 'pulled 2026-10-20'],
    ['Qualitative evidence (Stage 4)', 'FactSet UnstructuredContent (transcripts/news)', 'max 2 queries per company', 'pulled 2026-10-20'],
    ['Investment approach', 'Notion: Investment_Approach_and_Process (Part 1)', 'see intl/approach_summary.md', 'last edited 2026-05-19'],
]
sheet(wb, 'Sources', ['Data item', 'Source', 'Detail', 'As-of'], src, {'Data item': 28, 'Source': 50, 'Detail': 90})
out = os.path.join(ROOT, 'Buffett_Test_Data_Europe_India.xlsx'); wb.save(out); print('wrote', out, {w.title: w.max_row - 1 for w in wb.worksheets})
