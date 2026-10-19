"""Build Buffett_Test_Data.xlsx (Universe, Triage, Analysis, Sources) from stage files."""
import csv, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
from fsload import tkey
ROOT = os.path.join(os.path.dirname(__file__), '..')
D = os.path.join(ROOT, 'data')

s1 = list(csv.DictReader(open(os.path.join(D, 'stage1_universe.csv'))))
tri = {r['ticker']: r for r in csv.DictReader(open(os.path.join(ROOT, 'triage.csv')))}
res = {}
p = os.path.join(ROOT, 'analyses', '_results_final.jsonl')
if not os.path.exists(p): p = os.path.join(ROOT, 'analyses', '_results.jsonl')
for line in open(p):
    line = line.strip()
    if line:
        r = json.loads(line); res[r['ticker']] = r   # last line wins

def summ(t):
    f = os.path.join(D, tkey(t) + '_summary.json')
    return json.load(open(f)) if os.path.exists(f) else {}
def dat(t):
    f = os.path.join(D, tkey(t) + '.json')
    return json.load(open(f)) if os.path.exists(f) else {}

HDR = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='1F3864')
def sheet(wb, name, header, rows, widths=None):
    ws = wb.create_sheet(name); ws.append(header)
    for c in ws[1]: c.font = HDR; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical='top')
    for r in rows: ws.append(r)
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(18, len(h) + 2)))
    return ws

wb = Workbook(); wb.remove(wb.active)
# Universe: all Test-1 passes (+ near-misses, labelled)
rows = []
for r in s1:
    if r['test1'] not in ('PASS', 'NEAR_MISS'): continue
    s = summ(r['ticker']); d = dat(r['ticker'])
    rows.append([r['ticker'], r['name'], r['test1'], float(r['ntm_pe']) if r['ntm_pe'] else None, float(r['price_usd']) if r['price_usd'] else None,
                 float(r['ntm_eps_mean']) if r['ntm_eps_mean'] else None, r['earnings_basis'], r['is_reit'], d.get('market_value_usd_m'),
                 (d.get('entity') or {}).get('rbics_sector'), (d.get('entity') or {}).get('rbics_subsector'), r['source_note']])
rows.sort(key=lambda x: (x[2] != 'PASS', x[3] if x[3] is not None else 99))
sheet(wb, 'Universe', ['Ticker', 'Name', 'Test 1', 'NTM P/E', 'Price USD (19-Oct-26)', 'NTM EPS/FFO mean', 'Basis', 'REIT', 'Mkt cap USD m', 'RBICS sector', 'RBICS subsector', 'Source'], rows,
      {'Name': 34, 'RBICS sector': 22, 'RBICS subsector': 26, 'Source': 60})
# Triage
rows = [[t, r['name'], float(r['ntm_pe']) if r['ntm_pe'] else None, r['verdict'], r['reason'], r['review_override'], r['flags']] for t, r in tri.items()]
order = {'KEEP': 0, 'TOO HARD': 1, 'NO': 2}
rows.sort(key=lambda x: (order.get(x[3], 3), x[2] or 99))
ws = sheet(wb, 'Triage', ['Ticker', 'Name', 'NTM P/E', 'Verdict', 'Reason', 'Reviewer override', 'Flags'], rows, {'Name': 30, 'Reason': 90, 'Reviewer override': 60})
for row in ws.iter_rows(min_row=2):
    row[4].alignment = Alignment(wrap_text=True, vertical='top'); row[5].alignment = Alignment(wrap_text=True, vertical='top')
# Analysis
cols = ['ticker', 'name', 'triage', 'ntm_pe', 'mkt_cap_bn', 'eps_cagr_10y', 'cons_growth_fy1', 'cons_growth_fy2', 'cons_growth_fy3', 'driver_sum_pct', 'top_drivers',
        'p_test2', 'test2', 'p_test3', 'test3', 'tests_passed', 'near_miss', 'specialness', 'confidence', 'earnings_quality_flag', 'key_uncertainty', 'calibration_note']
rows = []
for t, r in res.items():
    s = summ(t)
    r.setdefault('tests_passed', 1 + (r.get('test2') == 'PASS') + (r.get('test3') == 'PASS'))
    extra = [s.get('years_eps_growth_ge7'), s.get('fcf_conversion_10y'), s.get('net_debt_ebitda'), s.get('roic_10y_avg'), s.get('roe_10y_avg'), s.get('adj_vs_gaap_gap_5y_avg_pct')]
    rows.append([r.get(c) for c in cols] + extra)
rows.sort(key=lambda x: (-(x[15] or 0), -(x[11] or 0), x[3] or 99))
ws = sheet(wb, 'Analysis', [c.replace('_', ' ') for c in cols] + ['yrs EPS growth >=7% (of 10)', 'FCF conv 10y', 'Net debt/EBITDA', 'ROIC 10y avg', 'ROE 10y avg', 'Adj vs GAAP gap 5y %'], rows,
           {'name': 30, 'top drivers': 40, 'specialness': 40, 'key uncertainty': 70, 'earnings quality flag': 40, 'calibration note': 50})
for row in ws.iter_rows(min_row=2):
    for i in (11, 13):
        if isinstance(row[i].value, (int, float)): row[i].number_format = '0%'
# Sources
src = [
    ['Universe screen', 'FactSet CompanyScreener', 'FREF_ENTITY_COUNTRY_ISO2=US; ISON_SECONDARY=0; FF_MKT_VAL>1000 (USD m); 0<FE_VALUATION_PE_MEAN (estimatePeriod 2) <15', '2026-10-19'],
    ['Price', 'FactSet GlobalPrices (prices, close, USD)', 'Latest close; 3 HK listings use 16-Oct-2026', '2026-10-19'],
    ['NTM EPS / FFO', 'FactSet EstimatesConsensus consensus_rolling, periodicity NTMA, USD', 'mean used; FFO for REITs where available', '2026-10-19'],
    ['FY0–FY3 consensus EPS, SALES, FFO; EPS LTG', 'FactSet EstimatesConsensus consensus_rolling ANN, USD', 'relativePeriod 0=last actual, 1..3=FY1..FY3; mean/median/std/high/low/count', '2026-10-19'],
    ['10-yr fundamentals', 'FactSet Fundamentals ANN, USD (FY2015–FY2026)', 'FF_SALES, FF_EBIT_OPER, FF_EBIT_OPER_MGN, FF_EBITDA_OPER, FF_NET_INC, FF_EPS_DIL, FF_EPS_DIL_BEF_UNUSUAL, FF_COM_SHS_OUT_EPS_DIL, FF_ROIC, FF_ROE, FF_FREE_CF, FF_OPER_CF, FF_CAPEX, FF_NET_DEBT, FF_DEBT_EBITDA_OPER, FF_DIV_CF, FF_STK_PURCH_CF, FF_ACQ_BUS_CF; audit-card formula stored per value in data/<TICKER>.json', 'pulled 2026-10-19'],
    ['Adjusted EPS history', 'FactSet EstimatesConsensus surprise (actual = surpriseAfter), ANN, USD', 'Street/adjusted basis; FFO actuals for REITs', 'pulled 2026-10-19'],
    ['Business description, RBICS', 'FactSet EntityReference', '', 'pulled 2026-10-19'],
    ['Market cap', 'FactSet GlobalPrices market_value (USD m); FMP batch-market-cap where FactSet timed out', 'see sources.market_value in data/<TICKER>.json', '2026-10-16/19'],
    ['Qualitative evidence (Stage 4)', 'FactSet UnstructuredContent (transcripts/news, Jul-2025–Oct-2026)', 'max 2 queries per company', 'pulled 2026-10-19'],
    ['Backup', 'Financial Modeling Prep (FMP)', 'used only where FactSet data missing; flagged FMP_BACKUP', '2026-10-19'],
    ['Investment approach', 'Notion: Investment_Approach_and_Process (Part 1)', 'summarised in approach_summary.md', 'last edited 2026-05-19'],
]
sheet(wb, 'Sources', ['Data item', 'Source', 'Detail', 'As-of'], src, {'Data item': 30, 'Source': 50, 'Detail': 90})
out = os.path.join(ROOT, 'Buffett_Test_Data.xlsx'); wb.save(out); print('wrote', out, {ws.title: ws.max_row - 1 for ws in wb.worksheets})
